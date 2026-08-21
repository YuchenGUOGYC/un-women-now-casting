from __future__ import annotations

import argparse
import csv
import re
from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any


DEFAULT_RADAR_CSV_DIR = Path("data/nmc_xinan_radar/roi_pipeline/csv")
DEFAULT_LOOKUP = Path("data/nmc_xinan_radar/roi_pipeline/lookup/radar_highres_to_lonlat_lookup.xlsx")
DEFAULT_POINT_OUTPUT_DIR = Path("data/nmc_xinan_radar/roi_pipeline/hourly/points_by_hour")
DEFAULT_SUMMARY_OUTPUT = Path("data/nmc_xinan_radar/roi_pipeline/hourly/hourly_weather_summary.csv")

TIMESTAMP_RE = re.compile(r"_(\d{14})\d{3}_")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Estimate hourly radar precipitation for 64 coarser points from "
            "per-image ROI radar echo CSVs. Echo high-resolution radar pixels are "
            "mapped to one of the 64 points through the lookup table. Each 6-minute "
            "pixel dBZ is converted to rainfall with a Z-R curve, accumulated over "
            "Beijing-time hours from UTC filename timestamps, then each coarse "
            "point uses the maximum of its assigned high-resolution pixels."
        )
    )
    parser.add_argument(
        "--radar-csv-dir",
        type=Path,
        default=DEFAULT_RADAR_CSV_DIR,
        help=f"Directory containing per-image *_roi_pixels.csv files. Default: {DEFAULT_RADAR_CSV_DIR}",
    )
    parser.add_argument(
        "--lookup",
        type=Path,
        default=DEFAULT_LOOKUP,
        help=f"High-resolution-radar-to-coarse-point lookup table. Default: {DEFAULT_LOOKUP}",
    )
    parser.add_argument(
        "--point-output-dir",
        type=Path,
        default=DEFAULT_POINT_OUTPUT_DIR,
        help=(
            "Output directory for per-hour point CSVs. Each hour gets one CSV "
            f"with 64 coarse-point rows. Default: {DEFAULT_POINT_OUTPUT_DIR}"
        ),
    )
    parser.add_argument(
        "--summary-output",
        type=Path,
        default=DEFAULT_SUMMARY_OUTPUT,
        help=f"Output CSV with one row per hour. Default: {DEFAULT_SUMMARY_OUTPUT}",
    )
    parser.add_argument(
        "--zr-a",
        type=float,
        default=200.0,
        help="Z-R curve A coefficient for Z = A * R^b. Default: 200",
    )
    parser.add_argument(
        "--zr-b",
        type=float,
        default=1.6,
        help="Z-R curve b exponent for Z = A * R^b. Default: 1.6",
    )
    parser.add_argument(
        "--interval-minutes",
        type=float,
        default=6.0,
        help="Minutes represented by each radar image. Default: 6",
    )
    return parser.parse_args()


def radar_grid_uid(pixel_x: Any, pixel_y: Any) -> str:
    return f"grid_x{int(float(pixel_x))}_y{int(float(pixel_y))}"


def parse_image_utc_time(path: Path) -> datetime:
    match = TIMESTAMP_RE.search(path.name)
    if not match:
        raise ValueError(f"Cannot parse timestamp from filename: {path.name}")
    return datetime.strptime(match.group(1), "%Y%m%d%H%M%S")


def utc_to_beijing_time(value: datetime) -> datetime:
    return value + timedelta(hours=8)


def hour_key(value: datetime) -> str:
    return value.replace(minute=0, second=0, microsecond=0).strftime("%Y-%m-%d %H:00:00")


def read_csv_rows(path: Path) -> list[dict[str, Any]]:
    with path.open("r", newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def read_xlsx_rows(path: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to read .xlsx lookup files. Run: pip install -r requirements.txt") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    rows = worksheet.iter_rows(values_only=True)
    try:
        headers = [str(value).strip() if value is not None else "" for value in next(rows)]
    except StopIteration:
        return []

    records: list[dict[str, Any]] = []
    for row in rows:
        if not any(value is not None and str(value).strip() != "" for value in row):
            continue
        record = {}
        for index, header in enumerate(headers):
            if header:
                record[header] = row[index] if index < len(row) else None
        records.append(record)
    return records


def read_table(path: Path) -> list[dict[str, Any]]:
    if path.suffix.lower() == ".csv":
        return read_csv_rows(path)
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        return read_xlsx_rows(path)
    raise ValueError(f"Unsupported table format: {path.suffix}")


def clean_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def load_lookup(path: Path) -> tuple[dict[str, dict[str, Any]], dict[str, list[str]]]:
    rows = read_table(path)
    if not rows:
        raise RuntimeError(f"No lookup rows found: {path}")

    highres_to_point: dict[str, dict[str, Any]] = {}
    point_to_highres: dict[str, list[str]] = defaultdict(list)
    for row in rows:
        grid_uid = clean_text(row.get("radar_grid_uid"))
        point_uid = clean_text(row.get("matched_point_uid"))
        if not grid_uid or not point_uid:
            continue
        point_info = {
            "point_uid": point_uid,
            "point_lon": row.get("matched_point_lon", ""),
            "point_lat": row.get("matched_point_lat", ""),
        }
        for key, value in row.items():
            if str(key).startswith("matched_point_"):
                point_info[key] = value
        highres_to_point[grid_uid] = point_info
        point_to_highres[point_uid].append(grid_uid)

    if not point_to_highres:
        raise RuntimeError(f"No matched points found in lookup: {path}")
    return highres_to_point, dict(point_to_highres)


def to_float_or_none(value: Any) -> float | None:
    text = clean_text(value)
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def dbz_to_rain_rate_mm_per_hour(dbz: float, zr_a: float, zr_b: float) -> float:
    reflectivity_z = 10 ** (dbz / 10)
    return (reflectivity_z / zr_a) ** (1 / zr_b)


def dbz_to_interval_precip_mm(dbz: float, zr_a: float, zr_b: float, interval_minutes: float) -> float:
    return dbz_to_rain_rate_mm_per_hour(dbz, zr_a=zr_a, zr_b=zr_b) * interval_minutes / 60


def collect_hourly_echoes(
    radar_csv_dir: Path,
    highres_to_point: dict[str, dict[str, Any]],
    zr_a: float,
    zr_b: float,
    interval_minutes: float,
) -> tuple[dict[str, dict[str, dict[str, Any]]], dict[str, int]]:
    hourly_echoes: dict[str, dict[str, dict[str, Any]]] = defaultdict(
        lambda: defaultdict(lambda: {"echo_grids": set(), "dbz_values": [], "precip_by_grid": defaultdict(float)})
    )
    image_count_by_hour: dict[str, int] = defaultdict(int)

    csv_paths = sorted(radar_csv_dir.glob("*_roi_pixels.csv"))
    if not csv_paths:
        raise RuntimeError(f"No *_roi_pixels.csv files found in {radar_csv_dir}")

    for csv_path in csv_paths:
        image_hour = hour_key(utc_to_beijing_time(parse_image_utc_time(csv_path)))
        image_count_by_hour[image_hour] += 1
        with csv_path.open("r", newline="", encoding="utf-8-sig") as handle:
            for row in csv.DictReader(handle):
                dbz = to_float_or_none(row.get("dbz"))
                if dbz is None:
                    continue
                grid_uid = clean_text(row.get("radar_grid_uid"))
                if not grid_uid:
                    grid_uid = radar_grid_uid(row.get("pixel_x"), row.get("pixel_y"))
                point_info = highres_to_point.get(grid_uid)
                if point_info is None:
                    continue
                point_uid = point_info["point_uid"]
                interval_precip_mm = dbz_to_interval_precip_mm(
                    dbz,
                    zr_a=zr_a,
                    zr_b=zr_b,
                    interval_minutes=interval_minutes,
                )
                hourly_echoes[image_hour][point_uid]["echo_grids"].add(grid_uid)
                hourly_echoes[image_hour][point_uid]["dbz_values"].append(dbz)
                hourly_echoes[image_hour][point_uid]["precip_by_grid"][grid_uid] += interval_precip_mm
    return hourly_echoes, dict(image_count_by_hour)


def save_csv(path: Path, rows: list[dict[str, Any]], headers: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)


POINT_HEADERS = [
    "hour",
    "point_uid",
    "point_lon",
    "point_lat",
    "assigned_highres_grid_count",
    "echo_highres_grid_count",
    "echo_coverage",
    "max_dbz",
    "hourly_precip_mm",
    "max_precip_highres_grid_uid",
    "image_count_in_hour",
    "status",
]


SUMMARY_HEADERS = [
    "hour",
    "image_count_in_hour",
    "point_count",
    "echo_point_count",
    "point_level_clear_count",
    "point_level_precipitation_count",
    "max_hourly_precip_mm",
    "max_hourly_precip_point_uid",
    "max_hourly_precip_highres_grid_uid",
    "max_dbz",
    "max_dbz_point_uid",
    "hour_status",
    "total_echo_highres_grid_count",
]


def hour_output_name(hour: str) -> str:
    parsed = datetime.strptime(hour, "%Y-%m-%d %H:%M:%S")
    return f"hourly_point_weather_{parsed:%Y%m%d_%H%M}.csv"


def save_point_rows_by_hour(output_dir: Path, point_rows: list[dict[str, Any]]) -> list[Path]:
    rows_by_hour: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in point_rows:
        rows_by_hour[str(row["hour"])].append(row)

    output_paths: list[Path] = []
    for hour in sorted(rows_by_hour):
        path = output_dir / hour_output_name(hour)
        save_csv(path, rows_by_hour[hour], POINT_HEADERS)
        output_paths.append(path)
    return output_paths


def classify_hours(args: argparse.Namespace) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    highres_to_point, point_to_highres = load_lookup(args.lookup)
    hourly_echoes, image_count_by_hour = collect_hourly_echoes(
        radar_csv_dir=args.radar_csv_dir,
        highres_to_point=highres_to_point,
        zr_a=args.zr_a,
        zr_b=args.zr_b,
        interval_minutes=args.interval_minutes,
    )

    point_infos: dict[str, dict[str, Any]] = {}
    for info in highres_to_point.values():
        point_infos.setdefault(info["point_uid"], info)

    point_rows: list[dict[str, Any]] = []
    summary_rows: list[dict[str, Any]] = []
    for image_hour in sorted(image_count_by_hour):
        point_level_precipitation_count = 0
        point_level_clear_count = 0
        echo_point_count = 0
        total_echo_grids = 0
        max_hourly_precip_mm = 0.0
        max_hourly_precip_point_uid = ""
        max_hourly_precip_grid_uid = ""
        max_hour_dbz: float | str = ""
        max_hour_dbz_point_uid = ""

        for point_uid in sorted(point_to_highres):
            assigned_grids = point_to_highres[point_uid]
            point_hour = hourly_echoes.get(image_hour, {}).get(
                point_uid,
                {"echo_grids": set(), "dbz_values": [], "precip_by_grid": {}},
            )
            echo_count = len(point_hour["echo_grids"])
            dbz_values = point_hour["dbz_values"]
            precip_by_grid = point_hour["precip_by_grid"]
            max_dbz = max(dbz_values) if dbz_values else ""
            coverage = echo_count / len(assigned_grids) if assigned_grids else 0.0
            if precip_by_grid:
                max_grid_uid, point_precip_mm = max(precip_by_grid.items(), key=lambda item: item[1])
            else:
                max_grid_uid, point_precip_mm = "", 0.0
            point_precip_mm_rounded = round(point_precip_mm, 1)
            status = "precipitation" if point_precip_mm_rounded > 0 else "clear"

            if status == "precipitation":
                point_level_precipitation_count += 1
            else:
                point_level_clear_count += 1
            if echo_count > 0:
                echo_point_count += 1
            total_echo_grids += echo_count
            if point_precip_mm > max_hourly_precip_mm:
                max_hourly_precip_mm = point_precip_mm
                max_hourly_precip_point_uid = point_uid
                max_hourly_precip_grid_uid = max_grid_uid
            if dbz_values and (max_hour_dbz == "" or max(dbz_values) > max_hour_dbz):
                max_hour_dbz = max(dbz_values)
                max_hour_dbz_point_uid = point_uid

            info = point_infos.get(point_uid, {})
            point_rows.append(
                {
                    "hour": image_hour,
                    "point_uid": point_uid,
                    "point_lon": info.get("point_lon", ""),
                    "point_lat": info.get("point_lat", ""),
                    "assigned_highres_grid_count": len(assigned_grids),
                    "echo_highres_grid_count": echo_count,
                    "echo_coverage": round(coverage, 6),
                    "max_dbz": max_dbz,
                    "hourly_precip_mm": point_precip_mm_rounded,
                    "max_precip_highres_grid_uid": max_grid_uid,
                    "image_count_in_hour": image_count_by_hour[image_hour],
                    "status": status,
                }
            )

        point_count = len(point_to_highres)
        max_hourly_precip_mm_rounded = round(max_hourly_precip_mm, 1)
        hour_status = "precipitation" if max_hourly_precip_mm_rounded > 0 else "clear"
        summary_rows.append(
            {
                "hour": image_hour,
                "image_count_in_hour": image_count_by_hour[image_hour],
                "point_count": point_count,
                "echo_point_count": echo_point_count,
                "point_level_clear_count": point_level_clear_count,
                "point_level_precipitation_count": point_level_precipitation_count,
                "max_hourly_precip_mm": max_hourly_precip_mm_rounded,
                "max_hourly_precip_point_uid": max_hourly_precip_point_uid,
                "max_hourly_precip_highres_grid_uid": max_hourly_precip_grid_uid,
                "max_dbz": max_hour_dbz,
                "max_dbz_point_uid": max_hour_dbz_point_uid,
                "hour_status": hour_status,
                "total_echo_highres_grid_count": total_echo_grids,
            }
        )
    return point_rows, summary_rows


def main() -> None:
    args = parse_args()
    point_rows, summary_rows = classify_hours(args)
    point_paths = save_point_rows_by_hour(args.point_output_dir, point_rows)
    save_csv(args.summary_output, summary_rows, SUMMARY_HEADERS)
    print(f"Saved per-hour point CSVs: {args.point_output_dir} ({len(point_paths)} files, {len(point_rows)} rows)")
    print(f"Saved hourly summary: {args.summary_output} ({len(summary_rows)} rows)")


if __name__ == "__main__":
    main()
